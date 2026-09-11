"""Read-only: is the 20:00 emergency spike an energy problem or a power problem?

Reports SOC at the 4-hour boundaries of the exported schedule, then rebuilds the
per-slot residual load - PV - contract and checks it against the two playback
caps: stored energy above E_min, and the 5000 kW interface limit.
"""

from __future__ import annotations

import sys
from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl import load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parent))

from config import (
    ETA_CHARGE,
    ETA_DISCHARGE,
    EXP_DIR,
    E_MIN_KWH,
    FOUR_HOUR_BLOCKS,
    N_INTERVALS,
    P_MAX_KW,
    P_MAX_KWH,
    RESULT_DIR,
)
from export_results import template_header_slot
from load_data import load_prices, load_year_actuals


def main() -> int:
    labels = [label for _, _, label in FOUR_HOUR_BLOCKS]
    wb = load_workbook(RESULT_DIR / "result3.xlsx", data_only=True)
    ws = wb["充放电量"]

    days: dict[str, dict] = {}
    current = None
    for row in range(2, ws.max_row + 1):
        date_cell = ws.cell(row, 1).value
        if date_cell not in (None, ""):
            current = str(date_cell)
            days[current] = {"charge": {}, "discharge": {}, "soc0": None}
        if current is None:
            continue
        label = ws.cell(row, 2).value
        if label in labels:
            days[current]["charge"][label] = float(ws.cell(row, 3).value or 0.0)
            days[current]["discharge"][label] = float(ws.cell(row, 4).value or 0.0)
        if ws.cell(row, 5).value == "0:00":
            days[current]["soc0"] = float(ws.cell(row, 6).value)

    soc_at = {label: [] for label in labels}
    for payload in days.values():
        soc = payload["soc0"]
        for label in labels:
            soc += ETA_CHARGE * payload["charge"][label] - payload["discharge"][label] / ETA_DISCHARGE
            soc_at[label].append(soc)

    rows = []
    for label in labels:
        arr = np.asarray(soc_at[label])
        rows.append(
            {
                "clock": label.split("-")[1],
                "min": arr.min(),
                "p10": np.percentile(arr, 10),
                "median": np.median(arr),
                "p90": np.percentile(arr, 90),
                "days_near_floor": int((arr < E_MIN_KWH + 50).sum()),
                "n_days": arr.size,
            }
        )
    table = pd.DataFrame(rows)
    out_dir = EXP_DIR / "soc_profile"
    out_dir.mkdir(parents=True, exist_ok=True)
    table.to_csv(out_dir / "soc_at_block_boundaries.csv", index=False, encoding="utf-8-sig")
    pd.set_option("display.width", 200)
    print(table.round(1).to_string(index=False))

    # Per-slot residual against the playback caps.
    prices = load_prices()
    year = load_year_actuals(
        jan1_load=float(prices["typical_load_kw"].iloc[0]),
        jan1_pv=float(prices["typical_pv_kw"].iloc[0]),
    )
    dates = [str(pd.Timestamp(d).date()) for d in year["dates"]]
    idx = {d: i for i, d in enumerate(dates)}
    adj = wb["调整购电量"]
    header_slots = [
        template_header_slot(str(adj.cell(1, t + 2).value)) for t in range(N_INTERVALS)
    ]
    same_day = [(i, slot) for i, (off, slot) in enumerate(header_slots) if off == 0]

    res_rows = []
    for i, date_str in enumerate(days):
        row = i + 2
        d = idx[date_str]
        for col, slot in same_day:
            g = float(adj.cell(row, col + 2).value)
            residual = float(year["load_kwh"][d, slot] - year["pv_kwh"][d, slot] - g)
            res_rows.append({"hour": slot // 6, "residual": residual})
    res = pd.DataFrame(res_rows)

    print(f"\nper-slot power cap = {P_MAX_KWH:.2f} kWh/slot ({P_MAX_KW:.0f} kW)")
    print("residual = load - PV - contract, by hour (only hours that ever exceed the cap):")
    grouped = res.groupby("hour")["residual"]
    summary = pd.DataFrame(
        {
            "p50": grouped.median(),
            "p90": grouped.quantile(0.90),
            "max": grouped.max(),
            "slots_over_cap": grouped.apply(lambda s: int((s > P_MAX_KWH).sum())),
            "n_slots": grouped.size(),
        }
    )
    over = summary[summary["slots_over_cap"] > 0]
    print(over.round(1).to_string())
    summary.to_csv(out_dir / "residual_vs_power_cap.csv", encoding="utf-8-sig")

    # Exact per-slot attribution: replay the greedy rule and record, for every slot
    # that buys emergency energy, which of the two caps was the binding one.
    contract = np.zeros((len(days), N_INTERVALS))
    for i, date_str in enumerate(days):
        for col, slot in same_day:
            contract[i, slot] = float(adj.cell(i + 2, col + 2).value)
    # Slot 0 of each day sits in the previous row's tail column; day 0 is outside the window.
    tail_col = next(i for i, (off, _) in enumerate(header_slots) if off == 1)
    for i in range(1, len(days)):
        contract[i, 0] = float(adj.cell(i + 1, tail_col + 2).value)

    att = []
    for i, date_str in enumerate(days):
        d = idx[date_str]
        soc = days[date_str]["soc0"]
        for slot in range(N_INTERVALS):
            residual = float(
                year["load_kwh"][d, slot] - year["pv_kwh"][d, slot] - contract[i, slot]
            )
            if residual <= 0:
                charge = min(-residual, min(P_MAX_KWH, max(0.0, (10800.0 - soc) / ETA_CHARGE)))
                soc = min(10800.0, max(E_MIN_KWH, soc + ETA_CHARGE * charge))
                continue
            energy_cap = ETA_DISCHARGE * max(0.0, soc - E_MIN_KWH)
            cap = min(P_MAX_KWH, energy_cap)
            discharge = min(residual, cap)
            em = residual - discharge
            if em > 1e-6:
                att.append(
                    {
                        "hour": slot // 6,
                        "emergency_kwh": em,
                        "binding": "power" if P_MAX_KWH <= energy_cap else "energy",
                    }
                )
            soc = min(10800.0, max(E_MIN_KWH, soc - discharge / ETA_DISCHARGE))
    att = pd.DataFrame(att)
    total = att["emergency_kwh"].sum()
    print(f"\nreplayed emergency total {total:.1f} kWh (metrics says 82228.36)")
    print("\nwhich cap was binding:")
    print(att.groupby("binding")["emergency_kwh"].agg(["sum", "count"]).round(1).to_string())
    print("\nemergency by hour (exact per-slot), hours above 1% only:")
    by_hour = att.pivot_table(
        index="hour", columns="binding", values="emergency_kwh", aggfunc="sum"
    ).fillna(0.0)
    by_hour["total"] = by_hour.sum(axis=1)
    by_hour["pct"] = 100 * by_hour["total"] / total
    print(by_hour[by_hour["pct"] > 1].round(1).to_string())
    by_hour.to_csv(out_dir / "emergency_binding_cap.csv", encoding="utf-8-sig")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
