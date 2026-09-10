"""Write official result3.xlsx and specified-day paper tables."""

from __future__ import annotations

import shutil
from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl import load_workbook

from config import (
    ABS_TOL_KWH,
    FOUR_HOUR_BLOCKS,
    N_INTERVALS,
    PAPER_TABLE1_END_MINUTES,
    RESULT3_TEMPLATE_XLSX,
    RESULT_DIR,
    SPECIFIED_DATES,
)


def _block_sums(end_min: np.ndarray, values: np.ndarray) -> dict[str, float]:
    totals = {}
    for start, end, label in FOUR_HOUR_BLOCKS:
        mask = (end_min > start) & (end_min <= end)
        totals[label] = float(values[mask].sum())
    return totals


def merge_emergency(date_str: str, emergency: np.ndarray, end_min: np.ndarray) -> list[dict]:
    rows = []
    t = 0
    n = len(emergency)
    while t < n:
        if emergency[t] <= ABS_TOL_KWH:
            t += 1
            continue
        t0 = t
        qty = 0.0
        while t < n and emergency[t] > ABS_TOL_KWH:
            qty += float(emergency[t])
            t += 1
        start_m = int(end_min[t0] - 10)
        end_m = int(end_min[t - 1])
        rows.append(
            {
                "date": date_str,
                "period": f"{_label(start_m)}-{_label(end_m)}",
                "emergency_kwh": qty,
            }
        )
    return rows


def _label(minutes: int) -> str:
    if minutes >= 24 * 60:
        return "24:00"
    hour, minute = divmod(int(minutes), 60)
    return f"{hour}:{minute:02d}"


def export_result3(official: list[dict], end_min: np.ndarray, dest: Path | None = None) -> Path:
    dest = dest or (RESULT_DIR / "result3.xlsx")
    dest.parent.mkdir(parents=True, exist_ok=True)
    shutil.copy(RESULT3_TEMPLATE_XLSX, dest)
    wb = load_workbook(dest)
    plan_ws = wb["计划购电量"]
    adj_ws = wb["调整购电量"]
    cd_ws = wb["充放电量"]
    em_ws = wb["紧急购电量"]

    n_days = len(official)
    if plan_ws.max_row - 1 < n_days:
        raise ValueError(f"plan template rows {plan_ws.max_row - 1} < {n_days}")

    for i, rec in enumerate(official):
        row = i + 2
        plan_ws.cell(row, 1).value = rec["date"]
        adj_ws.cell(row, 1).value = rec["date"]
        g_plan = rec["g_plan_kwh"]
        g_adj = rec["g_adj_kwh"]
        for t in range(N_INTERVALS):
            plan_ws.cell(row, t + 2).value = float(g_plan[t])
            adj_ws.cell(row, t + 2).value = float(g_adj[t])
        plan_ws.cell(row, N_INTERVALS + 2).value = float(g_plan.sum())
        plan_ws.cell(row, N_INTERVALS + 3).value = float(rec["bill"]["plan_only_cost"])
        adj_ws.cell(row, N_INTERVALS + 2).value = float(g_adj.sum())
        adj_ws.cell(row, N_INTERVALS + 3).value = float(rec["bill"]["total_cost"])

    if cd_ws.max_row > 1:
        cd_ws.delete_rows(2, cd_ws.max_row - 1)
    row = 2
    for rec in official:
        charge_blocks = _block_sums(end_min, rec["actual"]["charge_kwh"])
        discharge_blocks = _block_sums(end_min, rec["actual"]["discharge_kwh"])
        for b, (_, _, label) in enumerate(FOUR_HOUR_BLOCKS):
            cd_ws.cell(row, 1).value = rec["date"] if b == 0 else None
            cd_ws.cell(row, 2).value = label
            cd_ws.cell(row, 3).value = charge_blocks[label]
            cd_ws.cell(row, 4).value = discharge_blocks[label]
            if b == 0:
                cd_ws.cell(row, 5).value = "0:00"
                cd_ws.cell(row, 6).value = float(rec["actual"]["soc0_kwh"])
            elif b == 1:
                cd_ws.cell(row, 5).value = "24:00"
                cd_ws.cell(row, 6).value = float(rec["actual"]["soc24_kwh"])
            row += 1

    if em_ws.max_row > 1:
        em_ws.delete_rows(2, em_ws.max_row - 1)
    row = 2
    for rec in official:
        merged = merge_emergency(rec["date"], rec["actual"]["emergency_kwh"], end_min)
        if not merged:
            em_ws.cell(row, 1).value = rec["date"]
            row += 1
            continue
        for j, item in enumerate(merged):
            em_ws.cell(row, 1).value = item["date"] if j == 0 else None
            em_ws.cell(row, 2).value = item["period"]
            em_ws.cell(row, 3).value = item["emergency_kwh"]
            row += 1

    wb.save(dest)
    return dest


def export_paper_tables(official: list[dict], end_min: np.ndarray, price144: np.ndarray) -> tuple[Path, Path, Path]:
    RESULT_DIR.mkdir(parents=True, exist_ok=True)
    by_date = {r["date"]: r for r in official}
    table1_rows = []
    table2_rows = []
    table3_rows = []
    for date_str in SPECIFIED_DATES:
        rec = by_date[date_str]
        g = rec["g_adj_kwh"]
        for end in PAPER_TABLE1_END_MINUTES:
            idx = int(np.where(end_min == end)[0][0])
            table1_rows.append(
                {
                    "date": date_str,
                    "end_min": int(end),
                    "purchase_kwh": float(g[idx]),
                }
            )
        table1_rows.append(
            {
                "date": date_str,
                "period": "all_day_kwh",
                "purchase_kwh": float(g.sum()),
            }
        )
        table1_rows.append(
            {
                "date": date_str,
                "period": "all_day_cost",
                "purchase_kwh": float(rec["bill"]["total_cost"]),
            }
        )
        charge_blocks = _block_sums(end_min, rec["actual"]["charge_kwh"])
        discharge_blocks = _block_sums(end_min, rec["actual"]["discharge_kwh"])
        for _, _, label in FOUR_HOUR_BLOCKS:
            table2_rows.append(
                {
                    "date": date_str,
                    "period": label,
                    "charge_kwh": charge_blocks[label],
                    "discharge_kwh": discharge_blocks[label],
                }
            )
        table2_rows.append(
            {
                "date": date_str,
                "period": "soc_0",
                "charge_kwh": rec["actual"]["soc0_kwh"],
                "discharge_kwh": rec["actual"]["soc24_kwh"],
            }
        )
        table3_rows.extend(merge_emergency(date_str, rec["actual"]["emergency_kwh"], end_min))

    p1 = RESULT_DIR / "table1_specified_days.csv"
    p2 = RESULT_DIR / "table2_specified_days.csv"
    p3 = RESULT_DIR / "table3_emergency.csv"
    pd.DataFrame(table1_rows).to_csv(p1, index=False, encoding="utf-8-sig")
    pd.DataFrame(table2_rows).to_csv(p2, index=False, encoding="utf-8-sig")
    pd.DataFrame(table3_rows).to_csv(p3, index=False, encoding="utf-8-sig")
    return p1, p2, p3
