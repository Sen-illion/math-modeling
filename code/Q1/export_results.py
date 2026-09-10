"""Write official result1.xlsx and paper Table 1 / Table 2 CSV files."""

from __future__ import annotations

import shutil
from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl import load_workbook

from config import (
    E0_KWH,
    FOUR_HOUR_BLOCKS,
    PAPER_TABLE1_END_MINUTES,
    RESULT1_TEMPLATE_XLSX,
    RESULT_DIR,
    ASSET_TABLE_DIR,
)


def _block_sums(frame: pd.DataFrame, values: np.ndarray) -> dict[str, float]:
    totals = {}
    for start, end, label in FOUR_HOUR_BLOCKS:
        mask = (frame["end_min"] > start) & (frame["end_min"] <= end)
        totals[label] = float(values[mask.to_numpy()].sum())
    return totals


def export_result1(frame: pd.DataFrame, dispatch: dict, dest: Path | None = None) -> Path:
    dest = dest or (RESULT_DIR / "result1.xlsx")
    dest.parent.mkdir(parents=True, exist_ok=True)
    shutil.copy(RESULT1_TEMPLATE_XLSX, dest)
    wb = load_workbook(dest)
    purchase_sheet = wb.worksheets[0]
    storage_sheet = wb.worksheets[1]

    purchase = np.asarray(dispatch["purchase_kwh"], dtype=float)
    if purchase_sheet.max_row - 1 != len(purchase):
        raise ValueError(
            f"template purchase rows {purchase_sheet.max_row - 1} != dispatch length {len(purchase)}"
        )
    for i, value in enumerate(purchase, start=2):
        idx = i - 2
        start_label = str(frame.loc[idx, "t_start"])
        end_label = str(frame.loc[idx, "t_end"])
        purchase_sheet.cell(i, 1).value = f"{start_label}-{end_label}"
        purchase_sheet.cell(i, 2).value = float(value)

    charge_blocks = _block_sums(frame, np.asarray(dispatch["charge_kwh"], dtype=float))
    discharge_blocks = _block_sums(frame, np.asarray(dispatch["discharge_kwh"], dtype=float))
    for row in range(2, 8):
        label = str(storage_sheet.cell(row, 1).value).strip()
        storage_sheet.cell(row, 2).value = charge_blocks[label]
        storage_sheet.cell(row, 3).value = discharge_blocks[label]

    # Template puts 0:00 SOC on row 2 col 5 and 24:00 SOC on row 3 col 5.
    storage_sheet.cell(2, 5).value = E0_KWH
    storage_sheet.cell(3, 5).value = float(dispatch["soc_end_kwh"][-1])
    wb.save(dest)
    return dest


def export_paper_tables(frame: pd.DataFrame, dispatch: dict, metrics: dict) -> tuple[Path, Path]:
    RESULT_DIR.mkdir(parents=True, exist_ok=True)
    ASSET_TABLE_DIR.mkdir(parents=True, exist_ok=True)

    rows = []
    for end_min in PAPER_TABLE1_END_MINUTES:
        match = frame.index[frame["end_min"] == end_min]
        if len(match) != 1:
            raise ValueError(f"cannot locate interval ending at {end_min} minutes")
        idx = int(match[0])
        start_label = frame.loc[idx, "t_start"]
        end_label = frame.loc[idx, "t_end"]
        rows.append(
            {
                "period": f"{start_label}-{end_label}",
                "purchase_kwh": float(dispatch["purchase_kwh"][idx]),
            }
        )
    table1 = pd.DataFrame(rows)
    table1.loc[len(table1)] = {
        "period": "daily_purchase_kwh",
        "purchase_kwh": float(np.asarray(dispatch["purchase_kwh"]).sum()),
    }
    table1.loc[len(table1)] = {
        "period": "daily_cost_yuan",
        "purchase_kwh": float(metrics["M0"]["cost"]),
    }

    charge_blocks = _block_sums(frame, np.asarray(dispatch["charge_kwh"], dtype=float))
    discharge_blocks = _block_sums(frame, np.asarray(dispatch["discharge_kwh"], dtype=float))
    table2_rows = [
        {
            "period": label,
            "charge_kwh": charge_blocks[label],
            "discharge_kwh": discharge_blocks[label],
        }
        for _, _, label in FOUR_HOUR_BLOCKS
    ]
    table2 = pd.DataFrame(table2_rows)
    table2.loc[len(table2)] = {"period": "soc_0:00_kwh", "charge_kwh": E0_KWH, "discharge_kwh": np.nan}
    table2.loc[len(table2)] = {
        "period": "soc_24:00_kwh",
        "charge_kwh": float(dispatch["soc_end_kwh"][-1]),
        "discharge_kwh": np.nan,
    }

    table1_path = RESULT_DIR / "table1.csv"
    table2_path = RESULT_DIR / "table2.csv"
    table1.to_csv(table1_path, index=False, encoding="utf-8-sig")
    table2.to_csv(table2_path, index=False, encoding="utf-8-sig")
    table1.to_csv(ASSET_TABLE_DIR / "Q1_table1.csv", index=False, encoding="utf-8-sig")
    table2.to_csv(ASSET_TABLE_DIR / "Q1_table2.csv", index=False, encoding="utf-8-sig")
    return table1_path, table2_path
