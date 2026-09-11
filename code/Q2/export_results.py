"""Write official result2.xlsx from calendar-aligned day-ahead plans."""

from __future__ import annotations

import shutil
import sys
from pathlib import Path

import numpy as np
from openpyxl import load_workbook

from config import (
    ABS_TOL_KWH,
    FOUR_HOUR_BLOCKS,
    N_INTERVALS,
    RESULT2_TEMPLATE_XLSX,
    RESULT_DIR,
)

_CODE_DIR = Path(__file__).resolve().parents[1]
if str(_CODE_DIR) not in sys.path:
    sys.path.insert(0, str(_CODE_DIR))

from common.time_slots import template_header_slot  # noqa: E402


def _block_sums(end_min: np.ndarray, values: np.ndarray) -> dict[str, float]:
    totals = {}
    for start, end, label in FOUR_HOUR_BLOCKS:
        mask = (end_min > start) & (end_min <= end)
        totals[label] = float(values[mask].sum())
    return totals


def _label(minutes: int) -> str:
    if minutes >= 24 * 60:
        return "24:00"
    hour, minute = divmod(int(minutes), 60)
    return f"{hour}:{minute:02d}"


def merge_emergency(date_str: str, emergency: np.ndarray, end_min: np.ndarray) -> list[dict]:
    rows = []
    t = 0
    n = len(emergency)
    while t < n:
        if emergency[t] <= ABS_TOL_KWH:
            t += 1
            continue
        t0 = t
        qty = 0.0
        while t < n and emergency[t] > ABS_TOL_KWH:
            qty += float(emergency[t])
            t += 1
        start_m = int(end_min[t0] - 10)
        end_m = int(end_min[t - 1])
        rows.append(
            {
                "date": date_str,
                "period": f"{_label(start_m)}-{_label(end_m)}",
                "emergency_kwh": qty,
            }
        )
    return rows


def _series_for_headers(official: list[dict], day_idx: int, headers: list[str], key: str) -> np.ndarray:
    rec = official[day_idx]
    nxt = official[day_idx + 1] if day_idx + 1 < len(official) else None
    out = np.zeros(len(headers), dtype=float)
    for i, header in enumerate(headers):
        day_off, slot = template_header_slot(header)
        src = rec if day_off == 0 else nxt
        if src is None:
            out[i] = 0.0
        else:
            out[i] = float(src[key][slot])
    return out


def export_result2(official: list[dict], end_min: np.ndarray, dest: Path | None = None) -> Path:
    dest = dest or (RESULT_DIR / "result2.xlsx")
    dest.parent.mkdir(parents=True, exist_ok=True)
    shutil.copy(RESULT2_TEMPLATE_XLSX, dest)
    wb = load_workbook(dest)
    plan_ws = wb["计划购电量"]
    cd_ws = wb["充放电量"]
    em_ws = wb["紧急购电量"]

    n_days = len(official)
    if plan_ws.max_row - 1 < n_days:
        raise ValueError(f"plan template rows {plan_ws.max_row - 1} < {n_days}")
    plan_headers = [plan_ws.cell(1, t + 2).value for t in range(N_INTERVALS)]
    for h in plan_headers:
        template_header_slot(str(h))

    for i, rec in enumerate(official):
        row = i + 2
        plan_ws.cell(row, 1).value = rec["date"]
        g_plan = _series_for_headers(official, i, plan_headers, "purchase_kwh")
        for t in range(N_INTERVALS):
            plan_ws.cell(row, t + 2).value = float(g_plan[t])
        plan_ws.cell(row, N_INTERVALS + 2).value = float(g_plan.sum())
        plan_ws.cell(row, N_INTERVALS + 3).value = float(rec["plan_cost"])

    if cd_ws.max_row > 1:
        cd_ws.delete_rows(2, cd_ws.max_row - 1)
    row = 2
    for rec in official:
        charge_blocks = _block_sums(end_min, rec["charge_kwh"])
        discharge_blocks = _block_sums(end_min, rec["discharge_kwh"])
        for b, (_, _, label) in enumerate(FOUR_HOUR_BLOCKS):
            cd_ws.cell(row, 1).value = rec["date"] if b == 0 else None
            cd_ws.cell(row, 2).value = label
            cd_ws.cell(row, 3).value = charge_blocks[label]
            cd_ws.cell(row, 4).value = discharge_blocks[label]
            if b == 0:
                cd_ws.cell(row, 5).value = "0:00"
                cd_ws.cell(row, 6).value = float(rec["soc0_kwh"])
            elif b == 1:
                cd_ws.cell(row, 5).value = "24:00"
                cd_ws.cell(row, 6).value = float(rec["soc24_kwh"])
            row += 1

    if em_ws.max_row > 1:
        em_ws.delete_rows(2, em_ws.max_row - 1)
    row = 2
    for rec in official:
        merged = merge_emergency(rec["date"], rec["emergency_kwh"], end_min)
        if not merged:
            em_ws.cell(row, 1).value = rec["date"]
            row += 1
            continue
        for j, item in enumerate(merged):
            em_ws.cell(row, 1).value = item["date"] if j == 0 else None
            em_ws.cell(row, 2).value = item["period"]
            em_ws.cell(row, 3).value = item["emergency_kwh"]
            row += 1

    wb.save(dest)
    return dest
