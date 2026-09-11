"""Export result4-2 / result4-3 without touching result2.xlsx or result3.xlsx."""

from __future__ import annotations

import json
import shutil
import sys
from datetime import datetime, timezone
from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl import load_workbook

from .config import (
    ABS_TOL_KWH,
    E0_FEB1_KWH,
    FOUR_HOUR_BLOCKS,
    N_INTERVALS,
    N_OFFICIAL_DAYS,
    OFFICIAL_END,
    OFFICIAL_START,
    PAPER_TABLE1_END_MINUTES,
    RESULT4_2_TEMPLATE,
    RESULT4_2_XLSX,
    RESULT4_3_TEMPLATE,
    RESULT4_3_XLSX,
    RESULT_DIR,
    SPECIFIED_DATES,
)

CODE_DIR = Path(__file__).resolve().parents[1]
if str(CODE_DIR) not in sys.path:
    sys.path.insert(0, str(CODE_DIR))

from common.time_slots import template_header_slot  # noqa: E402


def _unpadded_clock(minutes: int) -> str:
    if minutes >= 24 * 60:
        return "24:00"
    hour, minute = divmod(int(minutes), 60)
    return f"{hour}:{minute:02d}"


def merge_emergency_periods(emergency_kwh: np.ndarray, start_min: np.ndarray, end_min: np.ndarray) -> list[tuple[str, float]]:
    emergency = np.asarray(emergency_kwh, dtype=float)
    periods = []
    i = 0
    n = len(emergency)
    while i < n:
        if emergency[i] <= ABS_TOL_KWH:
            i += 1
            continue
        j = i
        total = 0.0
        while j < n and emergency[j] > ABS_TOL_KWH:
            total += float(emergency[j])
            j += 1
        label = f"{_unpadded_clock(int(start_min[i]))}-{_unpadded_clock(int(end_min[j - 1]))}"
        periods.append((label, total))
        i = j
    return periods


def _block_sums(end_min: np.ndarray, values: np.ndarray) -> dict[str, float]:
    totals = {}
    for start, end, label in FOUR_HOUR_BLOCKS:
        mask = (end_min > start) & (end_min <= end)
        totals[label] = float(np.asarray(values, dtype=float)[mask].sum())
    return totals


def _series_for_headers(official: list[dict], day_idx: int, headers: list[str], key: str) -> np.ndarray:
    rec = official[day_idx]
    nxt = official[day_idx + 1] if day_idx + 1 < len(official) else None
    out = np.zeros(len(headers), dtype=float)
    for i, header in enumerate(headers):
        day_off, slot = template_header_slot(str(header))
        src = rec if day_off == 0 else nxt
        out[i] = 0.0 if src is None else float(np.asarray(src[key], dtype=float)[slot])
    return out


def _official_q42(traces: list[dict]) -> list[dict]:
    start = pd.Timestamp(OFFICIAL_START)
    end = pd.Timestamp(OFFICIAL_END)
    return [row for row in traces if start <= pd.Timestamp(row["date"]) <= end]


def export_result4_2(prices: pd.DataFrame, traces: list[dict], dest: Path | None = None) -> Path:
    dest = dest or RESULT4_2_XLSX
    dest = dest.resolve()
    if dest.parent != RESULT_DIR.resolve():
        raise ValueError(f"Q4-2 export must stay under {RESULT_DIR}, got {dest}")
    dest.parent.mkdir(parents=True, exist_ok=True)
    shutil.copy(RESULT4_2_TEMPLATE, dest)
    official = _official_q42(traces)
    if len(official) != N_OFFICIAL_DAYS:
        raise ValueError(f"expected {N_OFFICIAL_DAYS} official days, got {len(official)}")
    end_min = prices["end_min"].to_numpy(dtype=int)
    start_min = prices["start_min"].to_numpy(dtype=int)
    wb = load_workbook(dest)
    purchase_sheet = wb["计划购电量"]
    storage_sheet = wb["充放电量"]
    emergency_sheet = wb["紧急购电量"]
    n_time_cols = N_INTERVALS
    plan_headers = [purchase_sheet.cell(1, t + 2).value for t in range(n_time_cols)]
    for header in plan_headers:
        template_header_slot(str(header))
    for day_i, row in enumerate(official):
        row_idx = day_i + 2
        stamp = pd.Timestamp(row["date"]).to_pydatetime().replace(hour=0, minute=0, second=0, microsecond=0)
        purchase_sheet.cell(row_idx, 1).value = stamp
        mapped = _series_for_headers(official, day_i, plan_headers, "purchase_kwh")
        for col, value in enumerate(mapped, start=2):
            purchase_sheet.cell(row_idx, col).value = float(value)
        purchase_sheet.cell(row_idx, n_time_cols + 2).value = float(mapped.sum())
        purchase_sheet.cell(row_idx, n_time_cols + 3).value = float(row["plan_cost"])
    if storage_sheet.max_row > 1:
        storage_sheet.delete_rows(2, storage_sheet.max_row - 1)
    storage_row = 2
    prev_soc24 = None
    for day_i, row in enumerate(official):
        stamp = pd.Timestamp(row["date"]).to_pydatetime().replace(hour=0, minute=0, second=0, microsecond=0)
        soc0 = float(row["soc0_kwh"])
        soc24 = float(row["soc24_kwh"])
        if day_i == 0 and abs(soc0 - E0_FEB1_KWH) > 1e-6:
            raise ValueError(f"Feb 1 00:00 SOC must be {E0_FEB1_KWH}, got {soc0}")
        if prev_soc24 is not None and abs(soc0 - prev_soc24) > 1e-6:
            raise ValueError(f"{row['date']} SOC0 != previous SOC24")
        prev_soc24 = soc24
        charge_blocks = _block_sums(end_min, row["charge_kwh"])
        discharge_blocks = _block_sums(end_min, row["discharge_kwh"])
        for block_i, (_, _, label) in enumerate(FOUR_HOUR_BLOCKS):
            storage_sheet.cell(storage_row, 1).value = stamp if block_i == 0 else None
            storage_sheet.cell(storage_row, 2).value = label
            storage_sheet.cell(storage_row, 3).value = charge_blocks[label]
            storage_sheet.cell(storage_row, 4).value = discharge_blocks[label]
            if block_i == 0:
                storage_sheet.cell(storage_row, 5).value = "0:00"
                storage_sheet.cell(storage_row, 6).value = soc0
            elif block_i == 1:
                storage_sheet.cell(storage_row, 5).value = "24:00"
                storage_sheet.cell(storage_row, 6).value = soc24
            storage_row += 1
    if emergency_sheet.max_row > 1:
        emergency_sheet.delete_rows(2, emergency_sheet.max_row - 1)
    emergency_row = 2
    for row in official:
        periods = merge_emergency_periods(row["emergency_kwh"], start_min, end_min)
        if not periods:
            continue
        stamp = pd.Timestamp(row["date"]).to_pydatetime().replace(hour=0, minute=0, second=0, microsecond=0)
        for period_i, (label, amount) in enumerate(periods):
            emergency_sheet.cell(emergency_row, 1).value = stamp if period_i == 0 else None
            emergency_sheet.cell(emergency_row, 2).value = label
            emergency_sheet.cell(emergency_row, 3).value = float(amount)
            emergency_row += 1
    wb.save(dest)
    return dest


def export_specified_q42(prices: pd.DataFrame, traces: list[dict]) -> list[Path]:
    RESULT_DIR.mkdir(parents=True, exist_ok=True)
    by_date = {str(pd.Timestamp(row["date"]).date()): row for row in _official_q42(traces)}
    end_min = prices["end_min"].to_numpy(dtype=int)
    start_min = prices["start_min"].to_numpy(dtype=int)
    table1, table2, table3 = [], [], []
    for date in SPECIFIED_DATES:
        row = by_date[date]
        g = np.asarray(row["purchase_kwh"], dtype=float)
        for end in PAPER_TABLE1_END_MINUTES:
            idx = int(np.where(end_min == end)[0][0])
            table1.append({"date": date, "end_min": int(end), "purchase_kwh": float(g[idx])})
        table1.append({"date": date, "period": "all_day_kwh", "purchase_kwh": float(g.sum())})
        table1.append({"date": date, "period": "all_day_cost", "purchase_kwh": float(row["plan_cost"])})
        charge_blocks = _block_sums(end_min, row["charge_kwh"])
        discharge_blocks = _block_sums(end_min, row["discharge_kwh"])
        for _, _, label in FOUR_HOUR_BLOCKS:
            table2.append(
                {"date": date, "period": label, "charge_kwh": charge_blocks[label], "discharge_kwh": discharge_blocks[label]}
            )
        table2.append(
            {"date": date, "period": "soc_0_24", "charge_kwh": row["soc0_kwh"], "discharge_kwh": row["soc24_kwh"]}
        )
        for label, amount in merge_emergency_periods(row["emergency_kwh"], start_min, end_min):
            table3.append({"date": date, "period": label, "emergency_kwh": amount})
    paths = [
        RESULT_DIR / "q42_table1_specified_days.csv",
        RESULT_DIR / "q42_table2_specified_days.csv",
        RESULT_DIR / "q42_table3_emergency.csv",
    ]
    pd.DataFrame(table1).to_csv(paths[0], index=False, encoding="utf-8-sig")
    pd.DataFrame(table2).to_csv(paths[1], index=False, encoding="utf-8-sig")
    pd.DataFrame(table3).to_csv(paths[2], index=False, encoding="utf-8-sig")
    return paths


def export_result4_3(official: list[dict], end_min: np.ndarray, dest: Path | None = None) -> Path:
    dest = dest or RESULT4_3_XLSX
    dest = dest.resolve()
    if dest.parent != RESULT_DIR.resolve():
        raise ValueError(f"Q4-3 export must stay under {RESULT_DIR}, got {dest}")
    dest.parent.mkdir(parents=True, exist_ok=True)
    if len(official) != N_OFFICIAL_DAYS:
        raise ValueError(f"expected {N_OFFICIAL_DAYS} official days, got {len(official)}")
    shutil.copy(RESULT4_3_TEMPLATE, dest)
    wb = load_workbook(dest)
    plan_ws = wb["计划购电量"]
    adj_ws = wb["调整购电量"]
    cd_ws = wb["充放电量"]
    em_ws = wb["紧急购电量"]
    plan_headers = [plan_ws.cell(1, t + 2).value for t in range(N_INTERVALS)]
    for h in plan_headers:
        template_header_slot(str(h))
    start_min = np.asarray(end_min, dtype=int) - 10
    for i, rec in enumerate(official):
        row = i + 2
        plan_ws.cell(row, 1).value = rec["date"]
        adj_ws.cell(row, 1).value = rec["date"]
        g_plan = _series_for_headers(official, i, plan_headers, "g_plan_kwh")
        g_adj = _series_for_headers(official, i, plan_headers, "g_adj_kwh")
        for t in range(N_INTERVALS):
            plan_ws.cell(row, t + 2).value = float(g_plan[t])
            adj_ws.cell(row, t + 2).value = float(g_adj[t])
        plan_ws.cell(row, N_INTERVALS + 2).value = float(g_plan.sum())
        plan_ws.cell(row, N_INTERVALS + 3).value = float(rec["bill"]["plan_only_cost"])
        adj_ws.cell(row, N_INTERVALS + 2).value = float(g_adj.sum())
        adj_ws.cell(row, N_INTERVALS + 3).value = float(rec["bill"]["total_cost"])
    if cd_ws.max_row > 1:
        cd_ws.delete_rows(2, cd_ws.max_row - 1)
    row = 2
    for rec in official:
        charge_blocks = _block_sums(end_min, rec["actual"]["charge_kwh"])
        discharge_blocks = _block_sums(end_min, rec["actual"]["discharge_kwh"])
        for b, (_, _, label) in enumerate(FOUR_HOUR_BLOCKS):
            cd_ws.cell(row, 1).value = rec["date"] if b == 0 else None
            cd_ws.cell(row, 2).value = label
            cd_ws.cell(row, 3).value = charge_blocks[label]
            cd_ws.cell(row, 4).value = discharge_blocks[label]
            if b == 0:
                cd_ws.cell(row, 5).value = "0:00"
                cd_ws.cell(row, 6).value = float(rec["actual"]["soc0_kwh"])
            elif b == 1:
                cd_ws.cell(row, 5).value = "24:00"
                cd_ws.cell(row, 6).value = float(rec["actual"]["soc24_kwh"])
            row += 1
    if em_ws.max_row > 1:
        em_ws.delete_rows(2, em_ws.max_row - 1)
    row = 2
    for rec in official:
        merged = merge_emergency_periods(rec["actual"]["emergency_kwh"], start_min, end_min)
        if not merged:
            em_ws.cell(row, 1).value = rec["date"]
            row += 1
            continue
        for j, (period, qty) in enumerate(merged):
            em_ws.cell(row, 1).value = rec["date"] if j == 0 else None
            em_ws.cell(row, 2).value = period
            em_ws.cell(row, 3).value = qty
            row += 1
    wb.save(dest)
    return dest


def export_paper_q43(official: list[dict], end_min: np.ndarray) -> tuple[Path, Path, Path]:
    RESULT_DIR.mkdir(parents=True, exist_ok=True)
    by_date = {r["date"]: r for r in official}
    start_min = np.asarray(end_min, dtype=int) - 10
    table1, table2, table3 = [], [], []
    for date_str in SPECIFIED_DATES:
        rec = by_date[date_str]
        g = rec["g_adj_kwh"]
        for end in PAPER_TABLE1_END_MINUTES:
            idx = int(np.where(end_min == end)[0][0])
            table1.append({"date": date_str, "end_min": int(end), "purchase_kwh": float(g[idx])})
        table1.append({"date": date_str, "period": "all_day_kwh", "purchase_kwh": float(g.sum())})
        table1.append({"date": date_str, "period": "all_day_cost", "purchase_kwh": float(rec["bill"]["total_cost"])})
        charge_blocks = _block_sums(end_min, rec["actual"]["charge_kwh"])
        discharge_blocks = _block_sums(end_min, rec["actual"]["discharge_kwh"])
        for _, _, label in FOUR_HOUR_BLOCKS:
            table2.append(
                {
                    "date": date_str,
                    "period": label,
                    "charge_kwh": charge_blocks[label],
                    "discharge_kwh": discharge_blocks[label],
                }
            )
        table2.append(
            {
                "date": date_str,
                "period": "soc_0_24",
                "charge_kwh": rec["actual"]["soc0_kwh"],
                "discharge_kwh": rec["actual"]["soc24_kwh"],
            }
        )
        for period, qty in merge_emergency_periods(rec["actual"]["emergency_kwh"], start_min, end_min):
            table3.append({"date": date_str, "period": period, "emergency_kwh": qty})
    p1 = RESULT_DIR / "q43_table1_specified_days.csv"
    p2 = RESULT_DIR / "q43_table2_specified_days.csv"
    p3 = RESULT_DIR / "q43_table3_emergency.csv"
    pd.DataFrame(table1).to_csv(p1, index=False, encoding="utf-8-sig")
    pd.DataFrame(table2).to_csv(p2, index=False, encoding="utf-8-sig")
    pd.DataFrame(table3).to_csv(p3, index=False, encoding="utf-8-sig")
    return p1, p2, p3


def _jsonable(value):
    if isinstance(value, dict):
        return {str(k): _jsonable(v) for k, v in value.items()}
    if isinstance(value, (list, tuple)):
        return [_jsonable(v) for v in value]
    if isinstance(value, np.ndarray):
        return value.tolist()
    if isinstance(value, np.generic):
        return value.item()
    if isinstance(value, Path):
        return str(value)
    return value


def write_json(path: Path, payload: dict) -> Path:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(_jsonable(payload), ensure_ascii=False, indent=2), encoding="utf-8")
    return path


def merge_metrics(q42_summary: dict | None = None, q43_metrics: dict | None = None) -> dict:
    existing = {}
    path = RESULT_DIR / "metrics.json"
    if path.exists():
        existing = json.loads(path.read_text(encoding="utf-8"))
    if q42_summary is not None:
        existing["q4_2"] = q42_summary
    if q43_metrics is not None:
        existing["q4_3"] = q43_metrics
    existing["official_window"] = {"start": OFFICIAL_START, "end": OFFICIAL_END, "n_days": N_OFFICIAL_DAYS}
    existing["plan_price"] = "hat0_expanding_layers"
    existing["settle_price"] = "attachment4_actual"
    existing["q4_3_intraday_price"] = "locked_hat0"
    write_json(path, existing)
    return existing


def write_run_manifest(patch: dict) -> dict:
    path = RESULT_DIR / "run_manifest.json"
    existing = {}
    if path.exists():
        existing = json.loads(path.read_text(encoding="utf-8"))
    existing.update(_jsonable(patch))
    existing["created_at"] = datetime.now(timezone.utc).isoformat()
    existing.setdefault("official_window", {"start": OFFICIAL_START, "end": OFFICIAL_END, "n_days": N_OFFICIAL_DAYS})
    existing.setdefault("outputs", {})
    write_json(path, existing)
    return existing
